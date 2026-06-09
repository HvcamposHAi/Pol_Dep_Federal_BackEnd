"""Importação de leads a partir de CSV (;) ou Excel (.xlsx).

Decodifica CSV como UTF-8 com BOM (``utf-8-sig``); mapeia cabeçalhos via
``column_mapping`` (resiliente a corrupção ``�`` e a reordenação); normaliza
valores por tipo; deriva o telefone E.164 (chave natural) e faz upsert
fill-only-empty. Linhas sem telefone válido vão para ``rejected``.
"""

from __future__ import annotations

import csv
import io
from typing import Any

from openpyxl import load_workbook
from sqlalchemy.ext.asyncio import AsyncSession

from app.config import get_settings
from app.core import normalize
from app.core.column_mapping import (
    PHONE_COLUMN,
    PHONE_FALLBACK_COLUMN,
    resolve_header,
)
from app.repositories import lead_repo
from app.schemas.lead import ImportResult, RejectedRow

_COERCERS = {
    "text": normalize.clean_value,
    "int": normalize.parse_int,
    "numeric": normalize.parse_numeric,
    "bool": normalize.to_bool_sim_nao,
    "array": normalize.to_array,
    "date": normalize.parse_date,
    "datetime": normalize.parse_datetime,
    "cep": normalize.clean_cep,
    "e164": None,  # tratado à parte (precisa do país default)
}


def _coerce(kind: str, raw_value: Any) -> Any:
    if kind == "e164":
        return normalize.to_e164(raw_value, get_settings().default_phone_country)
    func = _COERCERS.get(kind, normalize.clean_value)
    return func(raw_value)


def _read_rows(filename: str, content: bytes) -> list[dict[str, str]]:
    """Lê o arquivo e devolve linhas como dict {cabeçalho_bruto: valor_string}."""
    name = (filename or "").lower()
    if name.endswith(".xlsx") or name.endswith(".xlsm"):
        return _read_xlsx(content)
    return _read_csv(content)


def _read_csv(content: bytes) -> list[dict[str, str]]:
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text), delimiter=";")
    rows = list(reader)
    if not rows:
        return []
    headers = rows[0]
    out: list[dict[str, str]] = []
    for raw in rows[1:]:
        if not any(cell.strip() for cell in raw):
            continue  # linha totalmente vazia
        out.append({headers[i]: (raw[i] if i < len(raw) else "") for i in range(len(headers))})
    return out


def _read_xlsx(content: bytes) -> list[dict[str, str]]:
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header_cells = next(rows_iter)
    except StopIteration:
        return []
    headers = [str(h) if h is not None else "" for h in header_cells]
    out: list[dict[str, str]] = []
    for raw in rows_iter:
        if raw is None or not any(c is not None and str(c).strip() for c in raw):
            continue
        out.append(
            {
                headers[i]: ("" if i >= len(raw) or raw[i] is None else str(raw[i]))
                for i in range(len(headers))
            }
        )
    wb.close()
    return out


def map_row(raw_row: dict[str, str]) -> tuple[dict[str, Any], str | None]:
    """Mapeia uma linha bruta para colunas do banco + raw_source.

    Retorna (mapped, reject_reason). ``reject_reason`` não-nulo => linha inválida.
    """
    mapped: dict[str, Any] = {}
    for header, value in raw_row.items():
        resolved = resolve_header(header)
        if resolved is None:
            continue  # cabeçalho desconhecido -> ignora
        column, kind = resolved
        coerced = _coerce(kind, value)
        if coerced is not None:
            mapped[column] = coerced

    # Telefone (chave natural): coluna E.164; fallback para a coluna bruta
    phone = mapped.get(PHONE_COLUMN)
    if not phone:
        fallback_raw = mapped.get(PHONE_FALLBACK_COLUMN)
        phone = normalize.to_e164(fallback_raw, get_settings().default_phone_country)
        if phone:
            mapped[PHONE_COLUMN] = phone

    if not phone:
        return mapped, "telefone ausente ou inválido (E.164)"

    # raw_source: linha original (auditoria/remapeamento)
    mapped["raw_source"] = {k: v for k, v in raw_row.items()}
    return mapped, None


async def import_bytes(
    session: AsyncSession,
    filename: str,
    content: bytes,
) -> ImportResult:
    settings = get_settings()
    if len(content) > settings.max_upload_bytes:
        return ImportResult(
            total=0,
            inserted=0,
            updated=0,
            skipped=0,
            errors=[f"arquivo excede o limite de {settings.max_upload_mb} MB"],
        )

    try:
        raw_rows = _read_rows(filename, content)
    except Exception as exc:  # noqa: BLE001 - reporta falha de parsing ao chamador
        return ImportResult(
            total=0, inserted=0, updated=0, skipped=0, errors=[f"falha ao ler arquivo: {exc}"]
        )

    rejected: list[RejectedRow] = []
    to_upsert: list[dict[str, Any]] = []
    for idx, raw in enumerate(raw_rows, start=2):  # linha 1 = cabeçalho
        mapped, reason = map_row(raw)
        if reason:
            rejected.append(
                RejectedRow(
                    linha=idx,
                    motivo=reason,
                    nome=mapped.get("nome_completo") or mapped.get("nome"),
                )
            )
            continue
        to_upsert.append(mapped)

    inserted = updated = 0
    batch = settings.enrich_batch_size
    for start in range(0, len(to_upsert), batch):
        chunk = to_upsert[start : start + batch]
        ins, upd = await lead_repo.upsert_many(session, chunk, provider="csv")
        inserted += ins
        updated += upd

    # skipped: linhas válidas que já existiam e não tinham nada a preencher
    skipped = max(0, len(to_upsert) - inserted - updated)
    return ImportResult(
        total=len(raw_rows),
        inserted=inserted,
        updated=updated,
        skipped=skipped,
        rejected=rejected,
    )
