"""E2E: endpoints /stats, /leads/field-coverage e /leads/export via ASGI.

Usa httpx.ASGITransport + override de get_db apontando para o SQLite de teste.
"""

from __future__ import annotations

import io
from datetime import date

import httpx
import pytest
from openpyxl import load_workbook

from app.db.session import get_db
from app.main import app
from app.repositories import lead_repo

pytestmark = pytest.mark.e2e

_SEED = [
    {
        "phone_e164": "+5541900000001",
        "nome_completo": "Ana Souza",
        "email": "ana@x.com",
        "cidade": "Curitiba",
        "estado": "PR",
        "data_nascimento": date(1990, 1, 1),
        "tags": ["vip", "apoiador"],
    },
    {"phone_e164": "+5541900000002", "nome_completo": "Bruno Dias", "estado": "SP"},
]


def _override(sm):
    async def _dep():
        async with sm() as s:
            yield s

    return _dep


@pytest.fixture
async def client(sessionmaker_fixture):
    sm = sessionmaker_fixture
    async with sm() as s:
        await lead_repo.upsert_many(s, _SEED)
        await s.commit()

    app.dependency_overrides[get_db] = _override(sm)
    transport = httpx.ASGITransport(app=app)
    async with httpx.AsyncClient(transport=transport, base_url="http://test") as c:
        yield c
    app.dependency_overrides.clear()


async def test_stats_endpoint(client):
    r = await client.get("/stats")
    assert r.status_code == 200
    body = r.json()
    assert body["total_leads"] == 2
    assert body["pendentes_apollo"] == 2
    assert isinstance(body["completude_media_pct"], float)


async def test_field_coverage_endpoint(client):
    r = await client.get("/leads/field-coverage", params={"estado": "PR"})
    assert r.status_code == 200
    cov = r.json()
    assert isinstance(cov, list)
    email = next(c for c in cov if c["field"] == "email")
    assert email["filled"] == 1  # só Ana no filtro PR
    assert email["missing"] == 0


async def test_export_csv_endpoint(client):
    r = await client.get("/leads/export", params={"format": "csv"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith("text/csv")
    assert "contatos.csv" in r.headers["content-disposition"]
    content = r.content
    assert content.startswith("﻿".encode())  # BOM UTF-8
    text = content.decode("utf-8-sig")
    header = text.splitlines()[0]
    assert ";" in header and "E-mail" in header  # separador ; e label do import
    assert "ana@x.com" in text
    # filtro respeitado: só Ana no PR
    r2 = await client.get("/leads/export", params={"format": "csv", "estado": "PR"})
    body2 = r2.content.decode("utf-8-sig")
    assert "ana@x.com" in body2 and "Bruno Dias" not in body2


async def test_export_xlsx_endpoint(client):
    r = await client.get("/leads/export", params={"format": "xlsx"})
    assert r.status_code == 200
    assert r.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    assert ws.max_row == 3  # cabeçalho + 2 leads
